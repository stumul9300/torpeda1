"""Финансовые отчёты (этап 1) по детализации из поддельного WB."""

import io
import json
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text

from torpeda.db import init_db
from torpeda.db.models import ReportRow

PERIOD = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


@pytest.fixture
def loaded(api, factory):
    assert api.post("/api/wb/sync").status_code == 200
    api.put("/api/settings", json={"tax_pct": "0"})
    api.put("/api/products/418020031/cost", json={"cost": "700"})
    with factory() as s:
        common = {"account_id": 1, "report_id": 5, "rr_dt": date(2026, 10, 21), "date_from": date(2026, 10, 14)}
        s.add_all(
            [
                ReportRow(rrd_id=100, nm_id=418020031, supplier_oper_name="Штраф", penalty=Decimal(500),
                          bonus_type_name="Занижение габаритов", **common),
                ReportRow(rrd_id=101, nm_id=None, supplier_oper_name="Удержание", deduction=Decimal(1000),
                          bonus_type_name="Оказание услуг «ВБ.Продвижение»", **common),
                ReportRow(rrd_id=102, nm_id=418020031, supplier_oper_name="Продажа", doc_type_name="Продажа",
                          quantity=0, cashback_discount=Decimal(30), **common),
            ]
        )  # fmt: skip
        s.commit()
    return api


def report(api, name, **params):
    resp = api.get(f"/api/reports/{name}", params=PERIOD | params)
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_range(loaded):
    r = loaded.get("/api/reports/range").json()
    assert (r["date_from"], r["date_to"]) == ("2026-10-20", "2026-10-21")


def test_detail_by_nm(loaded):
    data = report(loaded, "detail", group="nm")
    row = next(r for r in data["rows"] if r["nm_id"] == 418020031)
    assert (row["sales_qty"], row["returns_qty"], row["net_qty"]) == (10, 1, 9)
    assert row["payout"] == 15300  # 10 × 1700 − 1700
    assert row["supplier_sales"] == pytest.approx(27664.2)
    assert row["logistics"] == pytest.approx(4347.82)
    assert row["penalties"] == 500
    assert row["loyalty_comp"] == 30
    assert row["cost"] == 6300
    # 15300 + 30 − 4347,82 − 1,2 − 500 − 6300
    assert row["profit"] == pytest.approx(4180.98)
    assert row["sa_name"] is None or isinstance(row["sa_name"], str)
    other = next(r for r in data["rows"] if r["nm_id"] == "без артикула")
    assert other["deductions"] == 1000
    assert data["totals"]["profit"] == pytest.approx(4180.98 - 1000)


def test_detail_groups(loaded):
    by_office = report(loaded, "detail", group="office")
    assert by_office["columns"][0]["header"] == "Склад"
    by_size = report(loaded, "detail", group="office_nm_size")
    assert [c["field"] for c in by_size["columns"][:3]] == ["office_name", "nm_id", "sa_name"]


def test_weeks(loaded):
    data = report(loaded, "weeks")
    assert len(data["rows"]) == 1
    w = data["rows"][0]
    assert w["report_id"] == "5"
    # к оплате: 15300 + 30 − 4347,82 − 1,2 − 500 − 1000
    assert w["to_pay"] == pytest.approx(9480.98)
    assert w["wb_to_pay"] == pytest.approx(10950.98)  # шапка из finance-api
    assert w["diff"] == pytest.approx(9480.98 - 10950.98)


def test_logistics(loaded):
    row = report(loaded, "logistics")["rows"][0]
    assert row["deliveries"] == 30 and row["return_deliveries"] == 20
    assert row["logistics_per_unit"] == 483.09  # 4347,82 / 9, округлено до копеек
    assert report(loaded, "logistics", group="office")["columns"][0]["header"] == "Склад"


def test_penalties_and_deductions(loaded):
    pen = report(loaded, "penalties")
    assert [r["reason"] for r in pen["rows"]] == ["Занижение габаритов"]
    assert pen["totals"]["amount"] == 500
    ded = report(loaded, "deductions")
    assert ded["totals"]["amount"] == 1000
    assert "Продвижение" in ded["rows"][0]["reason"]


def test_xlsx(loaded):
    resp = loaded.get("/api/reports/detail.xlsx", params=PERIOD)
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws.cell(1, 1).value == "Артикул WB"
    assert ws.cell(2, 1).value == "Итого"


def test_unknown_report(loaded):
    assert loaded.get("/api/reports/nope", params=PERIOD).status_code == 404


def test_sync_report_period(api, fake_wb):
    resp = api.post("/api/wb/sync-report", json={"date_from": "2026-01-01", "date_to": "2026-10-24"})
    assert resp.json() == {"rows": 13}
    req = [r for r in fake_wb.requests if "sales-reports/detailed" in r.url.path][-1]
    assert json.loads(req.content)["dateFrom"] == "2026-01-01"
    assert api.post("/api/wb/sync-report", json={"date_from": "2026-02-01", "date_to": "2026-01-01"}).status_code == 400


def test_old_db_gets_new_columns(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'old.db'}")
    with engine.begin() as c:
        c.execute(text("CREATE TABLE report_rows (account_id INTEGER, rrd_id INTEGER, nm_id INTEGER, "
                       "PRIMARY KEY (account_id, rrd_id))"))  # fmt: skip
        c.execute(text("INSERT INTO report_rows VALUES (1, 7, 42)"))
    init_db(engine)
    cols = {c["name"] for c in inspect(engine).get_columns("report_rows")}
    assert {"ts_name", "bonus_type_name", "raw", "rebill_logistic_cost"} <= cols
    with engine.connect() as c:
        assert c.execute(text("SELECT nm_id FROM report_rows")).scalar() == 42


def test_finance_api_fields_are_mapped(api, factory, fake_wb):
    api.post("/api/wb/sync-report", json={"date_from": "2026-10-01", "date_to": "2026-10-24"})
    assert any("/api/finance/v1/sales-reports/detailed" in r.url.path for r in fake_wb.requests)
    assert not any("reportDetailByPeriod" in r.url.path for r in fake_wb.requests)
    with factory() as s:
        log = s.get(ReportRow, (1, 12))
        assert log.delivery_rub == Decimal("4347.82") and log.supplier_oper_name == "Логистика"


def test_falls_back_to_statistics_without_finance_scope(api, factory, fake_wb):
    fake_wb.finance_enabled = False
    assert api.post("/api/wb/sync-report", json={"date_from": "2026-10-01", "date_to": "2026-10-24"}).json() == {
        "rows": 13
    }
    assert any("reportDetailByPeriod" in r.url.path for r in fake_wb.requests)
