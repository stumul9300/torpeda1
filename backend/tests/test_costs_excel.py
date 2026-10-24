"""Себестоимость: выгрузка в Excel и загрузка обратно."""

import io

from openpyxl import Workbook, load_workbook


def xlsx(*rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(api, data: bytes):
    return api.post("/api/costs/import", files={"file": ("costs.xlsx", data)})


def cost(api, nm_id=418020031):
    return next(r for r in api.get("/api/products").json() if r["nm_id"] == nm_id)["cost"]


def test_export_then_import_roundtrip(api):
    api.post("/api/wb/sync")
    api.put("/api/products/418020031/cost", json={"cost": "700"})
    resp = api.get("/api/costs.xlsx")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Артикул WB", "Артикул продавца", "Название", "Бренд", "Себестоимость")
    assert rows[1] == (418020031, "8447", "Поло женское", "Stilly", 700)

    # тот же файл обратно — ничего не меняется
    assert upload(api, resp.content).json() == {"updated": 0, "unchanged": 1, "not_found": [], "errors": []}


def test_import_by_wb_article(api):
    api.post("/api/wb/sync")
    data = xlsx(("Артикул WB", "Себестоимость"), (418020031, "812,50"), (999, 100), (418020031, None))
    assert upload(api, data).json() == {"updated": 1, "unchanged": 0, "not_found": ["999"], "errors": []}
    assert float(cost(api)) == 812.5


def test_import_by_vendor_code_with_title_rows(api):
    api.post("/api/wb/sync")
    data = xlsx(("Мои товары",), (), ("Артикул продавца", "Название", "Себестоимость, ₽"), (8447, "Поло", 650))
    assert upload(api, data).json()["updated"] == 1
    assert float(cost(api)) == 650


def test_import_bad_values(api):
    api.post("/api/wb/sync")
    data = xlsx(("Артикул WB", "Себестоимость"), (418020031, "дорого"), (418020031, -5))
    body = upload(api, data).json()
    assert body["updated"] == 0
    assert len(body["errors"]) == 2


def test_import_wrong_file(api):
    assert upload(api, b"not excel").status_code == 400
    resp = upload(api, xlsx(("Что-то", "Другое"), (1, 2)))
    assert resp.status_code == 400
    assert "Себестоимость" in resp.json()["detail"]
