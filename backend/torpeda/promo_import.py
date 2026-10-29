"""Акции из файлов кабинета WB.

API не отдаёт товары автоакций, а в кабинете у каждой акции есть выгрузка в Excel («Все товары подходящие
для акции…», «Товары для возврата в акцию…», «Товары для акции…»). Название акции берётся из имени файла.
Файлы одной акции объединяются, повторная загрузка заменяет прежние данные акции. Такие акции хранятся
под отрицательным номером, чтобы не пересекаться с номерами WB.
"""

from __future__ import annotations

import io
import re
import zlib
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from torpeda.db import DEFAULT_ACCOUNT_ID
from torpeda.db.models import Promotion, PromotionItem

ACC = DEFAULT_ACCOUNT_ID
PREFIXES = ("Все товары подходящие для акции_", "Товары для возврата в акцию_", "Товары для акции_")
AUTO_SUFFIX = " автоматические скидки"
H_NM, H_PLAN, H_RETAIL, H_DISCOUNT, H_IN = (
    "Артикул WB",
    "Плановая цена для акции",
    "Текущая розничная цена",
    "Загружаемая скидка для участия в акции",
    "Товар уже участвует в акции",
)


def promo_name(filename: str) -> tuple[str, str]:
    """(название акции, тип) по имени файла: «Все товары подходящие для акции_11.11_ хиты автоматические
    скидки_29.10.2026 11.05.17.xlsx» → («11.11_ хиты», «auto»)."""
    name = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    name = re.sub(r"\.xlsx$", "", name, flags=re.I)
    name = re.sub(r"_\d{2}\.\d{2}\.\d{4} \d{2}\.\d{2}\.\d{2}$", "", name)
    for prefix in PREFIXES:
        if name.startswith(prefix):
            name = name[len(prefix) :]
            break
    kind = "regular"
    if name.endswith(AUTO_SUFFIX):
        name, kind = name[: -len(AUTO_SUFFIX)], "auto"
    return name.strip(), kind


def promo_id(name: str) -> int:
    return -(zlib.crc32(name.encode()) % 2_000_000_000) - 1


def _dec(v) -> Decimal | None:
    if v in (None, ""):
        return None
    try:
        return Decimal(str(v).replace(",", ".").replace(" ", ""))
    except InvalidOperation:
        return None


def parse(content: bytes) -> list[dict]:
    """Строки файла: артикул WB, цена до скидки, скидка для вступления, участвует ли уже."""
    try:
        ws = load_workbook(io.BytesIO(content), data_only=True).worksheets[0]
    except Exception as e:  # noqa: BLE001 - любой нечитаемый файл — одна и та же понятная ошибка
        raise ValueError("не получилось открыть файл как Excel (.xlsx)") from e
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    col = {str(h).strip(): i for i, h in enumerate(header) if h}
    missing = [h for h in (H_NM, H_PLAN, H_RETAIL) if h not in col]
    if missing:
        raise ValueError("нет колонок: " + ", ".join(missing))
    out = []
    for r in rows:
        nm, plan, retail = _dec(r[col[H_NM]]), _dec(r[col[H_PLAN]]), _dec(r[col[H_RETAIL]])
        if nm is None or plan is None or not retail:
            continue
        disc = _dec(r[col[H_DISCOUNT]]) if H_DISCOUNT in col else None
        if disc is None:  # скидка, при которой цена на сайте не выше плановой
            disc = (Decimal(1) - plan / retail) * 100
            disc = min(disc.to_integral_value(rounding="ROUND_CEILING"), Decimal(99))
        out.append(
            {
                "nm_id": int(nm),
                "plan_price": retail,
                "plan_discount_pct": max(disc, Decimal(0)),
                "in_action": str(r[col[H_IN]] or "").strip().lower() == "да" if H_IN in col else False,
            }
        )
    return out


def import_files(s: Session, files: list[tuple[str, bytes]]) -> dict:
    """Загружает файлы акций. Возвращает итог по каждому файлу."""
    cleared: set[int] = set()
    report = []
    for filename, content in files:
        name, kind = promo_name(filename)
        try:
            items = parse(content)
        except ValueError as e:
            report.append({"file": filename, "error": str(e)})
            continue
        pid = promo_id(name)
        s.merge(Promotion(account_id=ACC, id=pid, name=name, type=kind, start_at=None, end_at=None))
        if pid not in cleared:  # повторная загрузка заменяет данные акции, файлы одной загрузки складываются
            s.execute(delete(PromotionItem).where(PromotionItem.account_id == ACC, PromotionItem.promotion_id == pid))
            cleared.add(pid)
        for it in items:
            s.merge(PromotionItem(account_id=ACC, promotion_id=pid, **it))
        report.append({"file": filename, "promotion": name, "type": kind, "items": len(items)})
    s.commit()
    return {"files": report, "promotions": len(cleared), "items": sum(r.get("items", 0) for r in report)}


def file_promo_names(s: Session) -> set[str]:
    return {n for (n,) in s.execute(select(Promotion.name).where(Promotion.account_id == ACC, Promotion.id < 0))}
