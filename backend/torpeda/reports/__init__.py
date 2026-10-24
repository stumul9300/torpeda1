"""Отчёты вкладки «Отчёты»: каждый возвращает таблицу (колонки + строки + итог).

Одна и та же таблица показывается в браузере и выгружается в Excel.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font


@dataclass(frozen=True)
class Column:
    field: str
    header: str
    kind: str = "money"  # text | int | num | money | pct | date
    editable: bool = False  # ячейку можно править в браузере (например, план)


@dataclass
class Table:
    title: str
    columns: list[Column]
    rows: list[dict]
    totals: dict | None = None
    notes: list[str] = field(default_factory=list)

    def to_json(self) -> dict:
        return {
            "title": self.title,
            "columns": [c.__dict__ for c in self.columns],
            "rows": [_jsonable(r) for r in self.rows],
            "totals": _jsonable(self.totals) if self.totals else None,
            "notes": self.notes,
        }

    def to_xlsx(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = self.title[:31]
        ws.append([c.header for c in self.columns])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        rows = ([self.totals] if self.totals else []) + self.rows
        for r in rows:
            ws.append([_cell(r.get(c.field), c.kind) for c in self.columns])
        if self.totals:
            for cell in ws[2]:
                cell.font = Font(bold=True)
        for i, c in enumerate(self.columns, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 30 if c.kind == "text" else 14
        ws.freeze_panes = "A3" if self.totals else "A2"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def _cell(v, kind: str):
    if isinstance(v, Decimal):
        return float(round(v, 2))
    return v


def _jsonable(r: dict) -> dict:
    out = {}
    for k, v in r.items():
        if isinstance(v, Decimal):
            out[k] = float(round(v, 2))
        elif hasattr(v, "isoformat"):
            out[k] = v.isoformat()
        else:
            out[k] = v
    return out
