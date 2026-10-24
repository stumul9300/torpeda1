"""Себестоимость в Excel: выгрузка как страховка и обратная загрузка.

Файл — обычная таблица: «Артикул WB», «Артикул продавца», «Название», «Бренд», «Себестоимость».
При загрузке товар ищется по артикулу WB, а если его колонки нет — по артикулу продавца.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from torpeda.db import DEFAULT_ACCOUNT_ID
from torpeda.db.models import CostPrice, Product
from torpeda.services import current_costs

ACC = DEFAULT_ACCOUNT_ID
HEADERS = ["Артикул WB", "Артикул продавца", "Название", "Бренд", "Себестоимость"]


def export_costs(s: Session) -> bytes:
    costs = current_costs(s, date.today())
    wb = Workbook()
    ws = wb.active
    ws.title = "Себестоимость"
    ws.append(HEADERS)
    for p in s.scalars(select(Product).where(Product.account_id == ACC).order_by(Product.nm_id)):
        cost = costs.get(p.nm_id)
        ws.append([p.nm_id, p.vendor_code, p.title, p.brand, float(cost) if cost is not None else None])
    for col, width in zip("ABCDE", (14, 18, 50, 18, 14), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _find_columns(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """Ищет строку заголовка в первых 10 строках: номер строки и номера нужных колонок."""
    for i, row in enumerate(rows[:10]):
        cols: dict[str, int] = {}
        for j, cell in enumerate(row):
            name = str(cell or "").strip().lower()
            if "себестоим" in name:
                cols["cost"] = j
            elif ("артикул" in name and "wb" in name) or name in ("nm_id", "nmid", "артикул вб"):
                cols["nm"] = j
            elif "артикул" in name and ("продав" in name or "поставщ" in name):
                cols["vendor"] = j
        if "cost" in cols and ("nm" in cols or "vendor" in cols):
            return i, cols
    raise ValueError("Не нашла колонки «Себестоимость» и «Артикул WB» (или «Артикул продавца»).")


def _cell(row: tuple, j: int | None) -> str:
    if j is None or j >= len(row) or row[j] is None:
        return ""
    value = row[j]
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel хранит артикул 8447 как 8447.0
    return str(value).strip()


def _to_decimal(value: str) -> Decimal | None:
    if not value:
        return None
    try:
        d = Decimal(value.replace(" ", "").replace(" ", "").replace(",", "."))
    except InvalidOperation as e:
        raise ValueError(f"не число: {value}") from e
    if d < 0:
        raise ValueError(f"отрицательная: {value}")
    return d


def import_costs(s: Session, data: bytes, today: date | None = None) -> dict:
    today = today or date.today()
    try:
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    except Exception as e:
        raise ValueError("Не получилось открыть файл. Нужен файл Excel в формате .xlsx.") from e
    rows = list(ws.iter_rows(values_only=True))
    header, cols = _find_columns(rows)

    products = list(s.scalars(select(Product).where(Product.account_id == ACC)))
    by_nm = {p.nm_id: p for p in products}
    by_vendor = {str(p.vendor_code).strip(): p for p in products if p.vendor_code}
    costs = current_costs(s, today)

    updated = unchanged = 0
    not_found: list[str] = []
    errors: list[str] = []
    for n, row in enumerate(rows[header + 1 :], start=header + 2):
        try:
            cost = _to_decimal(_cell(row, cols["cost"]))
        except ValueError as e:
            errors.append(f"строка {n}: {e}")
            continue
        if cost is None:
            continue
        nm_raw = _cell(row, cols.get("nm"))
        vendor_raw = _cell(row, cols.get("vendor"))
        product = None
        if nm_raw:
            try:
                product = by_nm.get(int(float(nm_raw)))
            except ValueError:
                pass
        if product is None and vendor_raw:
            product = by_vendor.get(vendor_raw)
        if product is None:
            not_found.append(nm_raw or vendor_raw or f"строка {n}")
            continue
        if costs.get(product.nm_id) == cost:
            unchanged += 1
            continue
        s.add(CostPrice(account_id=ACC, nm_id=product.nm_id, cost=cost, valid_from=today))
        costs[product.nm_id] = cost
        updated += 1
    s.commit()
    return {"updated": updated, "unchanged": unchanged, "not_found": not_found, "errors": errors}
